#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Worker: PDF escaneado → OCR (Stirling) → Word limpio (plantilla CEFOA 3 págs).
Sin firmas, sellos, logos ni CamScanner.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

import urllib.request

# Reuse layout helpers by executing generation inline (python-docx)
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table

FONT = "Arial"

DEFAULTS = {
    "lugar_fecha": "Caracas, 291200OCT25",
    "epocas": "215° Independencia, 166° Federación y 26° Revolución",
    "refiere": "DEEJB",
    "archivo": "52-033010225000",
    "serial": "7033",
    "dependencia": "GT. FORMACIÓN",
    "ciudadano_rango": "GENERAL DE BRIGADA",
    "ciudadano_nombre": "RAFAEL ERNESTO ACUÑA ACEVEDO",
    "ciudadano_cargo1": "JEFE DEL SERVICIO DE ARMAMENTO",
    "ciudadano_cargo2": "DEL EJÉRCITO BOLIVARIANO",
    "ciudadano_lugar": "Fuerte Conopoima San Juan, Guárico.",
    "asunto": "SOLICITUD",
    "ref": "CEFOA N° 45",
    "telefono": "04166427379",
    "remitente": "NELSON JAVIER BOSCÁN PULGAR",
    "remitente_rango": "GENERAL DE DIVISIÓN",
    "remitente_cargo": "DIRECTOR DE EDUCACIÓN DEL EJÉRCITO BOLIVARIANO",
    "resolucion": (
        "SEGÚN RESOLUCIÓN N° 051932 DE FECHA 29 DE JULIO DE 2023 EMANADO DEL "
        "MINISTERIO DEL PODER POPULAR PARA LA DEFENSA"
    ),
    "iniciales": "ODRM/ardd",
    "fecha_corta": "29/10/2025",
    "fecha_form": "29-10-25",
    "periodo": "I-2025",
    "lugar_form": "CARACAS, DTTO. CAPITAL",
    "unidad": "CURSO ESPECIAL DE FORMACIÓN DE OFICIALES ASIMILADOS",
    "item1_desc": "CARTUCHO CAL. 9X19 MM (BALA)",
    "item1_qty": "1.560",
    "item2_desc": "CARTUCHO 7,62 X39 MM AK-103 (BALA)",
    "item2_qty": "4.680",
    "elaborado_nombre": "LEIDY CONTRERAS ROBLES",
    "elaborado_rango": "MAYOR",
    "elaborado_cargo": "CMDTE DE CUERPO DEL CEFOA",
    "revisado_nombre": "ANÍBAL DE JESÚS ZERPA ABREU",
    "revisado_rango": "CORONEL",
    "revisado_cargo": "DIRECTOR DEL LIMILAYA Y DEL CEFOA",
    "fecha_relacion": "Caricuao, 29 de octubre de 2025",
    "titulo_relacion": (
        "RELACIÓN DE LA SOLICITUD DE MUNICIÓN DE ADIESTRAMIENTO QUE SERÁ UTILIZADO "
        "POR EL PERSONAL DEL CURSO ESPECIAL DE FORMACIÓN DE OFICIALES ASIMILADOS N° 45"
    ),
}


def font(run, size=9, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def clear_cell(cell):
    for p in cell.paragraphs:
        p.clear()


def vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:vAlign"):
            tcPr.remove(child)
    el = OxmlElement("w:vAlign")
    el.set(qn("w:val"), "center")
    tcPr.append(el)


def write_cell(cell, text, *, size=8, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    clear_cell(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if text:
        for i, line in enumerate(str(text).split("\n")):
            if i:
                p.add_run().add_break()
            r = p.add_run(line)
            font(r, size=size, bold=bold, italic=italic)
    vcenter(cell)
    return p


def add_line(cell, text, *, size=8, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    font(r, size=size, bold=bold, italic=italic)
    return p


def set_table_borders(table, sz=8):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def no_borders(table):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def fixed(table):
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag == qn("w:tblLayout"):
            tblPr.remove(child)
    tblPr.append(layout)


def cell_border(cell, sz=18, edges=("left", "top", "bottom")):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcBorders"):
            tcPr.remove(child)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def nest_table(cell, rows, cols, border_sz=12):
    tc = cell._tc
    for child in list(tc):
        tc.remove(child)
    tbl = OxmlElement("w:tbl")
    tblPr = OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(border_sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    tbl.append(tblPr)
    grid = OxmlElement("w:tblGrid")
    for _ in range(cols):
        grid.append(OxmlElement("w:gridCol"))
    tbl.append(grid)
    for _ in range(rows):
        tr = OxmlElement("w:tr")
        for _c in range(cols):
            tcell = OxmlElement("w:tc")
            tcell.append(OxmlElement("w:tcPr"))
            tcell.append(OxmlElement("w:p"))
            tr.append(tcell)
        tbl.append(tr)
    tc.append(tbl)
    tc.append(OxmlElement("w:p"))
    return Table(tbl, cell)


def center_para(doc, text, *, size=10, bold=True, italic=False, space_before=0, space_after=0, underline=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    font(r, size=size, bold=bold, italic=italic)
    r.font.underline = underline
    return p


def left_para(doc, text, *, size=10, bold=False, space_before=0, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    font(r, size=size, bold=bold)
    return p


def right_para(doc, text, *, size=10, bold=False, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    font(r, size=size, bold=bold)
    return p


# ---------- OCR / parse ----------
def stirling_ocr(pdf_path: Path, stirling: str) -> str:
    """OCR via Stirling; returns plain text (best effort)."""
    boundary = "----OficioLimpioBoundary7MA4YWxkTrZu0gW"
    data = pdf_path.read_bytes()
    parts = []

    def add_field(name: str, value: str):
        parts.append(
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="{name}"\r\n\r\n'
            f"{value}\r\n".encode()
        )

    parts.append(
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="fileInput"; filename="{pdf_path.name}"\r\n'
        f"Content-Type: application/pdf\r\n\r\n".encode()
        + data
        + b"\r\n"
    )
    add_field("languages", "spa")
    add_field("languages", "eng")
    add_field("ocrType", "force-ocr")
    add_field("ocrRenderType", "hocr")
    add_field("deskew", "true")
    add_field("clean", "true")
    add_field("removeImagesAfter", "true")
    add_field("sidecar", "true")
    body = b"".join(parts) + f"--{boundary}--\r\n".encode()

    req = urllib.request.Request(
        f"{stirling.rstrip('/')}/api/v1/misc/ocr-pdf",
        data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=300) as resp:
            raw = resp.read()
    except Exception as e:
        print(f"WARN Stirling OCR falló: {e}", file=sys.stderr)
        return ""

    # ZIP (pdf+txt) or PDF
    if raw[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            texts = []
            for name in zf.namelist():
                if name.lower().endswith(".txt"):
                    texts.append(zf.read(name).decode("utf-8", errors="ignore"))
            if texts:
                return "\n".join(texts)
            # fallback: extract from pdf inside zip with pymupdf if available
            for name in zf.namelist():
                if name.lower().endswith(".pdf"):
                    try:
                        import pymupdf

                        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
                            tmp.write(zf.read(name))
                            tmp_path = tmp.name
                        doc = pymupdf.open(tmp_path)
                        t = "\n".join(p.get_text("text") for p in doc)
                        os.unlink(tmp_path)
                        return t
                    except Exception:
                        pass
    if raw[:4] == b"%PDF":
        try:
            import pymupdf

            doc = pymupdf.open(stream=raw, filetype="pdf")
            return "\n".join(p.get_text("text") for p in doc)
        except Exception:
            return ""
    return raw.decode("utf-8", errors="ignore")


def clean_noise(text: str) -> str:
    text = re.sub(r"(?i)escaneado con camscanner", "", text)
    text = re.sub(r"(?i)\bRECIBIDO\b", "", text)
    return text


def parse_fields(ocr: str) -> dict:
    f = dict(DEFAULTS)
    if not ocr:
        return f
    t = clean_noise(ocr)

    m = re.search(r"Lugar,?\s*Fecha:\s*([^\n]+)", t, re.I)
    if m:
        f["lugar_fecha"] = re.sub(r"\s+", " ", m.group(1)).strip(" ._")

    m = re.search(r"(\d{2,3}°\s*Independencia[^\n]+)", t, re.I)
    if m:
        f["epocas"] = re.sub(r"\s+", " ", m.group(1)).strip()

    m = re.search(r"refi[eé]rase\s*a:\s*([A-Z0-9]+)", t, re.I)
    if m:
        f["refiere"] = m.group(1).upper()

    m = re.search(r"(52-\d{12})", t)
    if m:
        f["archivo"] = m.group(1)

    m = re.search(r"N[°ºo]\s*Serial[^\d]*(\d{3,6})", t, re.I)
    if m:
        f["serial"] = m.group(1)

    m = re.search(r"Dependencia:\s*([^\n]+)", t, re.I)
    if m:
        f["dependencia"] = re.sub(r"\s+", " ", m.group(1)).strip(" .")

    m = re.search(r"REF:\s*(CEFOA[^\n]+)", t, re.I)
    if m:
        f["ref"] = re.sub(r"\s+", " ", m.group(1)).strip()

    m = re.search(r"(0\d{10,11})", t)
    if m:
        f["telefono"] = m.group(1)

    m = re.search(r"FECHA:\s*([\d\-/\.]+)", t, re.I)
    if m:
        f["fecha_form"] = m.group(1).strip()

    m = re.search(r"PERIODO DE CAMPO:\s*([^\n]+)", t, re.I)
    if m:
        f["periodo"] = re.sub(r"\s+", " ", m.group(1)).strip(" .")

    # quantities near 9X19 / AK
    m = re.search(r"9\s*[xX×]\s*19[^\d]{0,40}([\d\.]+)", t)
    if m:
        f["item1_qty"] = m.group(1)
    m = re.search(r"AK[-\s]?103[^\d]{0,40}([\d\.]+)", t, re.I)
    if m:
        f["item2_qty"] = m.group(1)

    m = re.search(r"(\d{1,2}/\d{1,2}/\d{2,4})", t)
    if m:
        f["fecha_corta"] = m.group(1)

    return f


# ---------- Word generation ----------
def generate_docx(fields: dict, out: Path) -> None:
    f = fields
    doc = Document()
    s0 = doc.sections[0]
    s0.orientation = WD_ORIENT.PORTRAIT
    s0.page_width = Cm(21.0)
    s0.page_height = Cm(29.7)
    s0.top_margin = Cm(1.0)
    s0.bottom_margin = Cm(0.9)
    s0.left_margin = Cm(1.4)
    s0.right_margin = Cm(1.4)

    # P1 header
    head = doc.add_table(rows=1, cols=2)
    fixed(head)
    no_borders(head)
    set_col_widths(head, [10.4, 7.8])
    L, R = head.cell(0, 0), head.cell(0, 1)
    write_cell(L, "REPÚBLICA BOLIVARIANA DE VENEZUELA", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_line(L, "MINISTERIO DEL PODER POPULAR PARA LA DEFENSA", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_line(L, "EJÉRCITO BOLIVARIANO", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_line(L, "", size=4)
    add_line(L, "1. CGEB", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_line(L, "2. DEEJ", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    box = nest_table(R, 5, 2, border_sz=12)
    box.cell(0, 0).merge(box.cell(0, 1))
    write_cell(box.cell(0, 0), f"Lugar, Fecha: {f['lugar_fecha']}", size=8, bold=True)
    box.cell(1, 0).merge(box.cell(1, 1))
    write_cell(box.cell(1, 0), f["epocas"], size=7, bold=True)
    box.cell(2, 0).merge(box.cell(2, 1))
    write_cell(box.cell(2, 0), f"Al contestar refiérase a: {f['refiere']}", size=8, bold=True)
    write_cell(box.cell(3, 0), f"N° de Archivo\n{f['archivo']}", size=7, bold=True)
    write_cell(box.cell(3, 1), f"N° Serial\n{f['serial']}", size=7, bold=True)
    box.cell(4, 0).merge(box.cell(4, 1))
    write_cell(box.cell(4, 0), f"Dependencia: {f['dependencia']}", size=8, bold=True)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    mid = doc.add_table(rows=1, cols=2)
    fixed(mid)
    no_borders(mid)
    set_col_widths(mid, [12.2, 6.0])
    ml, mr = mid.cell(0, 0), mid.cell(0, 1)
    cell_border(ml, sz=18, edges=("left", "top", "bottom"))
    write_cell(ml, "CIUDADANO:", size=10, bold=True)
    add_line(ml, f["ciudadano_rango"], size=10, bold=True)
    add_line(ml, f["ciudadano_nombre"], size=10, bold=True)
    add_line(ml, f["ciudadano_cargo1"], size=9, bold=True)
    add_line(ml, f["ciudadano_cargo2"], size=9, bold=True)
    add_line(ml, "AC: SU DESPACHO.", size=9, bold=True)
    add_line(ml, f["ciudadano_lugar"], size=9)
    write_cell(mr, f"ASUNTO: {f['asunto']}", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_line(mr, "", size=6)
    add_line(mr, f"REF: {f['ref']}", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    body1 = (
        "Tengo el agrado de dirigirme a usted, en la oportunidad de hacerle llegar un cordial saludo "
        "Bolivariano y Chavista, en nombre del Personal Militar y Civil que labora en esta prestigiosa "
        "Dirección a mi cargo, y a su vez, solicitar de sus buenos oficios, en el sentido de que estudie "
        "la posibilidad de autorizar la dotación de Abastecimiento Clase V, según lista anexa, los cuales "
        "serán usados en el Ejercicio de Adiestramiento a efectuarse durante el periodo de campo que se "
        "impartirá a los Aspirantes a Oficiales que integran el Curso Especial de Formación de Oficiales "
        "en la Categoría de Asimilados, a efectuarse desde el 14 de noviembre hasta el 27 de noviembre "
        "del presente año, con la finalidad de adiestrar a los futuros Oficiales Asimilados del Glorioso "
        "Ejército Bolivariano."
    )
    phrases = [
        "Bolivariano y Chavista",
        "Personal Militar y Civil",
        "Abastecimiento Clase V",
        "Ejercicio de Adiestramiento",
        "Aspirantes a Oficiales que integran el Curso Especial de Formación de Oficiales en la Categoría de Asimilados",
        "Glorioso Ejército Bolivariano",
    ]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    cursor = 0
    while cursor < len(body1):
        hit = ph_hit = None
        for ph in phrases:
            idx = body1.find(ph, cursor)
            if idx >= 0 and (hit is None or idx < hit):
                hit, ph_hit = idx, ph
        if hit is None:
            r = p.add_run(body1[cursor:])
            font(r, 10)
            break
        if hit > cursor:
            r = p.add_run(body1[cursor:hit])
            font(r, 10)
        r = p.add_run(body1[hit : hit + len(ph_hit)])
        font(r, 10, bold=True)
        cursor = hit + len(ph_hit)

    body2 = (
        "Solicitud que hago a usted respetuosamente, para sus conocimientos y demás fines consiguientes. "
        f"Quedando a sus gratas órdenes por el N° Tlf {f['telefono']}."
    )
    p2 = doc.add_paragraph(body2)
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p2.paragraph_format.space_after = Pt(18)
    for r in p2.runs:
        font(r, 10)

    center_para(doc, f["remitente"], size=10, bold=True, space_before=6)
    center_para(doc, f["remitente_rango"], size=10, bold=True, italic=True)
    center_para(doc, f["remitente_cargo"], size=10, bold=True)
    center_para(doc, f["resolucion"], size=7, bold=True, space_before=4, space_after=8, underline=True)

    p = doc.add_paragraph(f["iniciales"])
    p.paragraph_format.space_after = Pt(0)
    font(p.runs[0], 8)
    p = doc.add_paragraph(f["fecha_corta"])
    p.paragraph_format.space_after = Pt(4)
    font(p.runs[0], 8)

    center_para(doc, "ACUSE DE RECIBO", size=9, bold=True, space_after=2)
    ac_wrap = doc.add_table(rows=1, cols=2)
    fixed(ac_wrap)
    no_borders(ac_wrap)
    set_col_widths(ac_wrap, [15.8, 2.4])
    ac = nest_table(ac_wrap.cell(0, 0), 4, 16, border_sz=6)
    for i, h in enumerate([""] + [str(n) for n in range(1, 13)] + ["D", "M", "A"]):
        write_cell(ac.cell(0, i), h, size=7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(ac.cell(1, 0), "HORA\nMañana", size=6, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(ac.cell(2, 0), "Tarde", size=6, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c in range(1, 16):
        write_cell(ac.cell(1, c), "", size=6)
        write_cell(ac.cell(2, c), "", size=6)
    for i, h in enumerate([""] + [str(n) for n in range(13, 25)] + ["", "", ""]):
        write_cell(ac.cell(3, i), h, size=7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(ac_wrap.cell(0, 1), "\n\n___________\nFIRMA", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    an = doc.add_paragraph("ANEXO: Lo Indicado")
    an.paragraph_format.space_before = Pt(6)
    font(an.runs[0], 9, bold=True)

    # P2 landscape
    s1 = doc.add_section()
    s1.orientation = WD_ORIENT.LANDSCAPE
    s1.page_width = Cm(29.7)
    s1.page_height = Cm(21.0)
    s1.top_margin = Cm(0.8)
    s1.bottom_margin = Cm(0.7)
    s1.left_margin = Cm(1.0)
    s1.right_margin = Cm(1.0)

    center_para(doc, "REPÚBLICA BOLIVARIANA DE VENEZUELA", size=11, bold=True, space_after=0)
    center_para(doc, "MINISTERIO DEL PODER POPULAR PARA LA DEFENSA", size=10, bold=True, space_after=0)
    center_para(doc, "EJÉRCITO BOLIVARIANO", size=10, bold=True, space_after=0)
    center_para(doc, "DIRECCIÓN DE EDUCACIÓN DEL EJÉRCITO BOLIVARIANO", size=9, bold=True, space_after=0)
    center_para(doc, f["unidad"], size=9, bold=True, space_after=4)

    meta = doc.add_table(rows=3, cols=4)
    set_table_borders(meta, 10)
    fixed(meta)
    set_col_widths(meta, [7.2, 7.2, 5.5, 6.5])
    meta.cell(0, 0).merge(meta.cell(1, 0))
    write_cell(meta.cell(0, 0), "SOLICITUD DE MUNICIÓN\nDE ADIESTRAMIENTO", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(meta.cell(0, 1), f"LUGAR: {f['lugar_form']}", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(meta.cell(0, 2), f"FECHA: {f['fecha_form']}", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(meta.cell(0, 3), f"PERIODO DE CAMPO: {f['periodo']}", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(meta.cell(1, 1), "PÁGINA N°: 01 DE 01", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    meta.cell(1, 2).merge(meta.cell(1, 3))
    write_cell(meta.cell(1, 2), "", size=8)
    meta.cell(2, 0).merge(meta.cell(2, 3))
    write_cell(meta.cell(2, 0), f"UNIDAD: {f['unidad']}", size=9, bold=True)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    cols = ["N°", "CÓDIGO", "DESCRIPCIÓN", "MUNICIÓN\nEN EXISTENCIA", "SALDO\nFINAL", "DÉFICIT", "SOLICITUD", "TOTAL\nGENERAL"]
    q1, q2 = f["item1_qty"], f["item2_qty"]
    data = [
        ["01", "---", f["item1_desc"], "---", "---", q1, q1, q1],
        ["02", "---", f["item2_desc"], "---", "---", q2, q2, q2],
        ["03", "---", "---", "---", "---", "---", "---", "---"],
        ["04", "---", "---", "---", "---", "---", "---", "---"],
        ["05", "---", "---", "---", "---", "---", "---", "---"],
        ["06", "---", "---", "---", "---", "---", "---", "---"],
    ]
    tbl = doc.add_table(rows=1 + len(data), cols=8)
    set_table_borders(tbl, 8)
    fixed(tbl)
    set_col_widths(tbl, [1.2, 2.2, 8.8, 3.2, 2.5, 2.5, 2.8, 2.8])
    for i, h in enumerate(cols):
        write_cell(tbl.cell(0, i), h, size=7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            al = WD_ALIGN_PARAGRAPH.LEFT if ci == 2 else WD_ALIGN_PARAGRAPH.CENTER
            write_cell(tbl.cell(ri + 1, ci), val, size=9, bold=True, align=al)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    sig = doc.add_table(rows=1, cols=3)
    set_table_borders(sig, 12)
    fixed(sig)
    set_col_widths(sig, [9.0, 9.0, 9.5])
    blocks = [
        ("ELABORADO:", ["", "", "_______________________________", f["elaborado_nombre"], f["elaborado_rango"], f["elaborado_cargo"]]),
        ("REVISADO:", ["", "", "_______________________________", f["revisado_nombre"], f["revisado_rango"], f["revisado_cargo"]]),
        ("V°B°:", ["", "", "_______________________________", f["remitente"], f["remitente_rango"], "DIRECTOR DE EDUCACIÓN DEL", "EJÉRCITO BOLIVARIANO"]),
    ]
    for i, (title, lines) in enumerate(blocks):
        cell = sig.cell(0, i)
        write_cell(cell, title, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        for ln in lines:
            add_line(cell, ln if ln else " ", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # P3 relación
    s2 = doc.add_section()
    s2.orientation = WD_ORIENT.PORTRAIT
    s2.page_width = Cm(21.0)
    s2.page_height = Cm(29.7)
    s2.top_margin = Cm(1.2)
    s2.bottom_margin = Cm(1.0)
    s2.left_margin = Cm(1.5)
    s2.right_margin = Cm(1.5)

    center_para(doc, "REPÚBLICA BOLIVARIANA DE VENEZUELA", size=10, bold=True, space_after=0)
    center_para(doc, "MINISTERIO DEL PODER POPULAR PARA LA DEFENSA", size=9, bold=True, space_after=0)
    center_para(doc, "EJÉRCITO BOLIVARIANO", size=9, bold=True, space_after=0)
    center_para(doc, "DIRECCIÓN DE EDUCACIÓN DEL EJÉRCITO BOLIVARIANO", size=9, bold=True, space_after=0)
    center_para(doc, "EJÉRCITO BOLIVARIANO", size=9, bold=True, space_after=6)
    right_para(doc, f["fecha_relacion"], size=10, bold=True, space_after=8)
    center_para(doc, f["titulo_relacion"], size=9, bold=True, space_after=10)

    left_para(doc, "A. SITUACIÓN ACTUAL DEL PERSONAL DE LA UNIDAD EN FECHA 29OCT25:", size=9, bold=True, space_after=4)
    ta = doc.add_table(rows=3, cols=6)
    set_table_borders(ta, 8)
    fixed(ta)
    set_col_widths(ta, [3.5, 2.4, 2.4, 2.0, 2.8, 2.4])
    for i, h in enumerate(["SITUACIÓN", "OFIC SUP", "OFIC SUB", "T/P", "ASP/OFIC", "TOTAL"]):
        write_cell(ta.cell(0, i), h, size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, v in enumerate(["PERSONAL", "05", "14", "22", "89", "130"]):
        write_cell(ta.cell(1, i), v, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    for i, v in enumerate(["DISPONIBLES", "05", "14", "22", "89", "130"]):
        write_cell(ta.cell(2, i), v, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)

    left_para(doc, "B. RELACIÓN DE MUNICIÓN DE 9MM EN BASE A LA RELACIÓN DE 130 COMBATIENTES", size=9, bold=True, space_before=10, space_after=4)
    tb = doc.add_table(rows=6, cols=5)
    set_table_borders(tb, 8)
    fixed(tb)
    set_col_widths(tb, [1.2, 5.5, 3.5, 2.8, 2.8])
    for i, h in enumerate(["N°", "EJERCICIO", "POSICIÓN", "CARTUCHOS", "SUB TOTAL"]):
        write_cell(tb.cell(0, i), h, size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tb.cell(1, 0), "1.", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tb.cell(1, 1), "FAMILIARIZACIÓN", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tb.cell(1, 2), "PIE", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tb.cell(1, 3), "3", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tb.cell(1, 4), "390", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    tb.cell(2, 0).merge(tb.cell(4, 0))
    tb.cell(2, 1).merge(tb.cell(4, 1))
    write_cell(tb.cell(2, 0), "2.", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tb.cell(2, 1), "25 MTS", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r, pos in enumerate(["PIE", "RODILLA", "TENDIDO"]):
        write_cell(tb.cell(2 + r, 2), pos, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(tb.cell(2 + r, 3), "3", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(tb.cell(2 + r, 4), "390", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    tb.cell(5, 0).merge(tb.cell(5, 3))
    write_cell(tb.cell(5, 0), "TOTAL", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tb.cell(5, 4), q1, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    left_para(doc, "C. RELACIÓN DE MUNICIÓN DE AK-103 EN BASE A LA RELACIÓN DE 130 COMBATIENTES", size=9, bold=True, space_before=10, space_after=4)
    tc = doc.add_table(rows=10, cols=5)
    set_table_borders(tc, 8)
    fixed(tc)
    set_col_widths(tc, [1.2, 5.5, 3.5, 2.8, 2.8])
    for i, h in enumerate(["N°", "EJERCICIO", "POSICIÓN", "CARTUCHOS", "SUB TOTAL"]):
        write_cell(tc.cell(0, i), h, size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(1, 0), "1.", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(1, 1), "FAMILIARIZACIÓN", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(1, 2), "PIE", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(1, 3), "3", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(1, 4), "390", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(2, 0), "2.", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(2, 1), "TRIANGULACIÓN", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(2, 2), "PIE", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(2, 3), "3", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(2, 4), "390", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    tc.cell(3, 0).merge(tc.cell(5, 0))
    tc.cell(3, 1).merge(tc.cell(5, 1))
    write_cell(tc.cell(3, 0), "3.", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(3, 1), "75 MTS DIURNO", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r, pos in enumerate(["PIE", "RODILLA", "TENDIDO"]):
        write_cell(tc.cell(3 + r, 2), pos, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(tc.cell(3 + r, 3), "5", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(tc.cell(3 + r, 4), "650", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    tc.cell(6, 0).merge(tc.cell(8, 0))
    tc.cell(6, 1).merge(tc.cell(8, 1))
    write_cell(tc.cell(6, 0), "4.", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(6, 1), "25 MTS NOCTURNO", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r, pos in enumerate(["PIE", "RODILLA", "TENDIDO"]):
        write_cell(tc.cell(6 + r, 2), pos, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(tc.cell(6 + r, 3), "5", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(tc.cell(6 + r, 4), "650", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    tc.cell(9, 0).merge(tc.cell(9, 3))
    write_cell(tc.cell(9, 0), "TOTAL", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(tc.cell(9, 4), q2, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    center_para(doc, f"CNEL. {f['revisado_nombre']}", size=10, bold=True, space_before=18)
    center_para(doc, "DIRECTOR DEL CURSO ESPECIAL DE FORMACIÓN DE OFICIALES ASIMILADOS", size=9, bold=True)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pdf", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--stirling", default=os.getenv("STIRLING_URL", "https://pdf.klowshop.com"))
    args = ap.parse_args()

    pdf = Path(args.pdf)
    out = Path(args.out)
    if not pdf.is_file():
        print(f"PDF no encontrado: {pdf}", file=sys.stderr)
        sys.exit(2)

    print(f"OCR via {args.stirling} …", flush=True)
    ocr = stirling_ocr(pdf, args.stirling)
    fields = parse_fields(ocr)
    print(f"Campos OCR: {len(ocr)} chars", flush=True)
    generate_docx(fields, out)
    print(f"OK {out} ({out.stat().st_size} bytes)", flush=True)
    print("META_JSON:" + json.dumps({"ocr_chars": len(ocr), "archivo": fields.get("archivo"), "serial": fields.get("serial")}), flush=True)


if __name__ == "__main__":
    main()
